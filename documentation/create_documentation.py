from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(__file__).resolve().parent


def report_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=29, alignment=TA_CENTER, textColor=colors.HexColor("#16324F"), spaceAfter=18))
    styles.add(ParagraphStyle(name="ReportSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=12, leading=16, alignment=TA_CENTER, textColor=colors.HexColor("#35607D"), spaceAfter=12))
    styles.add(ParagraphStyle(name="Chapter", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#16324F"), spaceBefore=8, spaceAfter=10, keepWithNext=True))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#237A77"), spaceBefore=10, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name="BodyReport", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=14, alignment=TA_LEFT, textColor=colors.HexColor("#24313A"), spaceAfter=6))
    styles.add(ParagraphStyle(name="SmallReport", parent=styles["BodyText"], fontName="Helvetica", fontSize=8, leading=11, textColor=colors.HexColor("#394A54"), spaceAfter=3))
    styles.add(ParagraphStyle(name="BulletReport", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, leftIndent=14, firstLineIndent=-8, bulletIndent=2, textColor=colors.HexColor("#24313A"), spaceAfter=3))
    return styles


def p(text, style):
    return Paragraph(text, style)


def bullet(text, style):
    return Paragraph("&#8226; " + text, style)


def table(data, widths=None):
    converted = []
    for row in data:
        converted.append([cell if hasattr(cell, "wrap") else Paragraph(str(cell), report_styles()["SmallReport"]) for cell in row])
    result = Table(converted, colWidths=widths, repeatRows=1)
    result.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C7CF")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F2F7F8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F2F7F8"), colors.white]),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return result


def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#B7C7CF"))
    canvas.line(0.65 * inch, 0.55 * inch, A4[0] - 0.65 * inch, 0.55 * inch)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#5B6B73"))
    canvas.drawString(0.65 * inch, 0.35 * inch, "Dungeon Exploration | Project Documentation")
    canvas.drawRightString(A4[0] - 0.65 * inch, 0.35 * inch, f"Page {doc.page}")
    canvas.restoreState()


def report_story():
    s = report_styles()
    story = [
        Spacer(1, 1.0 * inch),
        p("A Project Report", s["ReportSubtitle"]),
        p("Dungeon Exploration", s["ReportTitle"]),
        p("A 2D procedural action-adventure game built with Python and Pygame", s["ReportSubtitle"]),
        Spacer(1, 0.35 * inch),
        p("Prepared from the Dungeon Exploration project source, the documentation sample, and the supplied Dungeon Quest / Dungeon Exploration reports.", s["BodyReport"]),
        Spacer(1, 1.2 * inch),
        p("September 2026", s["ReportSubtitle"]),
        PageBreak(),
        p("Abstract", s["Chapter"]),
        p("Dungeon Exploration is a desktop 2D action-adventure game in which the player creates or selects an explorer profile, chooses a class, enters a procedurally generated dungeon, fights monsters, collects keys and treasures, uses consumable items, and either clears the level or records a failed run. The project uses Python, Pygame, JSON profile persistence, grid-based navigation, and a hybrid procedural generation pipeline. The AI portion is primarily a decision and generation system: a genetic algorithm evolves candidate layouts, cellular automata smooth them, breadth-first search measures reachability and repairs disconnected regions, A* validates paths, and a decision-tree-style profile analysis adjusts future monster counts. This report follows the structure of the supplied sample documentation while documenting the implementation that is actually present in this repository.", s["BodyReport"]),
        p("Acknowledgement", s["Chapter"]),
        p("This report was assembled from the project source and the supplied reference documents. It is intended to support project evaluation and presentation of the current implementation.", s["BodyReport"]),
        p("List of Abbreviations", s["Chapter"]),
        table([["Abbreviation", "Meaning"], ["AI", "Artificial Intelligence"], ["BFS", "Breadth-First Search"], ["A*", "A-star pathfinding"], ["JSON", "JavaScript Object Notation"], ["UI", "User Interface"], ["FPS", "Frames Per Second"]], [1.4 * inch, 4.6 * inch]),
        PageBreak(),
        p("Contents", s["Chapter"]),
        p("1 Introduction", s["Section"]),
        p("2 Literature Review", s["Section"]),
        p("3 Methodology", s["Section"]),
        p("4 System Analysis", s["Section"]),
        p("5 System Design and Internal Flow", s["Section"]),
        p("6 AI and Procedural Generation", s["Section"]),
        p("7 System Development and Implementation", s["Section"]),
        p("8 Testing and Debugging", s["Section"]),
        p("9 Conclusion and Future Work", s["Section"]),
        PageBreak(),
        p("1 Introduction", s["Chapter"]),
        p("1.1 Overview", s["Section"]),
        p("The game is a single-process Pygame application with a screen-manager architecture. The player starts at a setup screen, enters a menu, and begins a run. Each run receives a fresh random seed and a new world generated from the active profile. The dungeon is a grid of tiles containing walls, floor, water, lava, traps, treasures, keys, scrolls, lights, decorations, and monsters.", s["BodyReport"]),
        p("1.2 Problem Statement", s["Section"]),
        p("A fixed dungeon quickly becomes predictable. The project addresses this by generating a different playable layout for each run while preserving a coherent objective: explore, collect the door key, survive enemies, and reach the locked door. The game also needs to remember progression without requiring a server or database.", s["BodyReport"]),
        p("1.3 Objectives", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["Provide a playable 2D dungeon exploration loop.", "Generate connected and varied dungeons at runtime.", "Support profile creation, class selection, inventory, progression, and persistent statistics.", "Use pathfinding and reachability checks to keep entities and objectives accessible.", "Adapt future dungeon difficulty using recent player outcomes."]],
        p("1.4 Features", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["Explorer profile creation and local JSON persistence.", "Screen flow for setup, menu, game, inventory, profile, pause, and completion/death states.", "Keyboard movement, sprinting, terrain speed effects, mouse aiming, projectiles, and monster combat.", "Keys, treasures, scrolls, coins, bow/light rewards, traps, water, lava, and locked-door progression.", "Scheduled monster spawning and run-history recording."]],
        p("1.5 Significance", s["Section"]),
        p("The project demonstrates how game programming, procedural content generation, search algorithms, persistence, UI state management, and adaptive difficulty can be combined into one small but complete interactive system.", s["BodyReport"]),
        p("1.6 Scope and Limitations", s["Section"]),
        p("The current implementation is a local desktop game. The ClientSession name is retained from the project design, but persistence is file-based through `src/client/data/profiles.json`; there is no multiplayer server, SQL database, or online account service. The AI is rule-based and algorithmic rather than a trained machine-learning model. Monster movement and combat behavior are implemented in the game screen/entity layer, while the most significant AI contribution is map generation and profile-aware difficulty selection.", s["BodyReport"]),
        PageBreak(),
        p("2 Literature Review", s["Chapter"]),
        p("2.1 Shattered Pixel Dungeon", s["Section"]),
        p("This reference demonstrates the replay value of procedural dungeon levels, class choices, resource management, and grid-based exploration. Dungeon Exploration adopts the replayability principle while using real-time Pygame controls and a custom hybrid generator.", s["BodyReport"]),
        p("2.2 Elona+", s["Section"]),
        p("Elona+ illustrates the depth possible through classes, exploration, items, and emergent combinations. The current project uses a smaller scope: class statistics, inventory rewards, consumable effects, and persistent profile progression.", s["BodyReport"]),
        p("2.3 Multiplayer Dungeon Quest", s["Section"]),
        p("The supplied report identifies static maps and limited item variety as common limitations. This project responds with runtime map generation, traps, liquids, multiple reward types, scheduled enemies, and profile-aware map scoring. It remains a local single-player implementation.", s["BodyReport"]),
        p("2.4 Related Game Design Lessons", s["Section"]),
        table([["Reference idea", "Applied project response"], ["Procedural replayability", "Genetic layout population, crossover, mutation, and randomized carving"], ["Resource pressure", "Health, stamina, terrain effects, consumables, keys, and coins"], ["Readable objectives", "Farthest-reachable door key, locked door, guarded treasure"], ["Adaptive challenge", "Recent-game rules adjust the monster limit"]], [2.0 * inch, 4.0 * inch]),
        PageBreak(),
        p("3 Methodology", s["Chapter"]),
        p("3.1 Software Development Cycle", s["Section"]),
        p("An iterative SDLC approach fits this project because gameplay features can be implemented and tested in small slices: screen flow, profile persistence, player movement, generation, entity placement, combat, and progression. The supplied sample documentation also recommends an iterative model for evolving requirements and early playable prototypes.", s["BodyReport"]),
        p("3.2 Technologies and Tools", s["Section"]),
        table([["Technology", "Use in the project"], ["Python 3", "Application logic, data structures, algorithms, and persistence"], ["Pygame 2.5.8", "Window, input, rendering, timing, sprites, and audio-ready game loop"], ["JSON", "Local player profiles and progression storage"], ["VS Code", "Development and debugging environment"], ["ReportLab / python-pptx", "Generation of this deliverable" ]], [1.55 * inch, 4.45 * inch]),
        p("3.3 Main Modules", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["`main.py`: adds `src/` and `src/client/` to the import path and starts the application.", "`src/client/main.py`: initializes Pygame, registers screens, owns the main loop, and shuts down persistence.", "`src/client/session.py`: loads/saves profiles, registers/selects explorers, records runs, and commits persistent inventory.", "`src/shared/mapgen/`: contains constants, layout, genetic algorithm, BFS, A*, placement, decision logic, and world generation.", "`src/client/screens/game_screen.py`: connects generated worlds to player movement, interactions, combat, spawning, and outcomes."]],
        PageBreak(),
        p("4 System Analysis", s["Chapter"]),
        p("4.1 Requirement Analysis", s["Section"]),
        table([["ID", "Functional requirement", "Implementation evidence"], ["FR1", "Create or select an explorer profile", "SetupScreen and ClientSession"], ["FR2", "Navigate through a menu and select a class", "MenuScreen, profile state, config classes"], ["FR3", "Generate a playable dungeon", "shared.mapgen.generator.generate_map"], ["FR4", "Move, sprint, fight, and interact", "GameScreen and Player"], ["FR5", "Collect items and unlock progression", "placement, inventory, and world interaction methods"], ["FR6", "Persist progression and recent outcomes", "profiles.json and ClientSession"], ["FR7", "Adapt future challenge", "decision_tree.monster_limit_for_profile"]], [0.55 * inch, 2.5 * inch, 2.95 * inch]),
        p("4.2 Non-functional Requirements", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["Playability: generation repairs connectivity and validates paths before placing important entities.", "Maintainability: map generation is separated into focused modules instead of one monolithic function.", "Responsiveness: the game loop uses a clock delta and forwards update/draw/event calls to one active screen.", "Reproducibility: generation accepts a seed, allowing a world to be recreated for debugging.", "Portability: the project is Python-based and uses local assets and JSON rather than a required external service."]],
        p("4.3 Feasibility", s["Section"]),
        p("Technical feasibility is high because the project uses a small dependency set and standard grid algorithms. Operational feasibility is supported by direct controls and a clear menu-to-game flow. Economic feasibility is favorable for an academic project because the software stack is open source and the game uses local storage. Schedule risk is concentrated in art assets, balancing, and edge cases in gameplay rather than in external infrastructure.", s["BodyReport"]),
        PageBreak(),
        p("5 System Design and Internal Flow", s["Chapter"]),
        p("5.1 Overall Game Flow", s["Section"]),
        p("Launch -> initialize Pygame and ClientSession -> SetupScreen creates/selects profile -> MenuScreen -> Start Game -> GameScreen generates a seeded world -> player explores and fights -> collect key and objective items -> unlock door and clear level OR health reaches zero -> record run outcome -> retain persistent progress -> return to menu or restart.", s["BodyReport"]),
        p("5.2 Runtime Loop", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["`clock.tick(FPS)` produces frame delta time.", "Pygame events are sent to the active screen.", "The active screen updates player state, world spawn timers, monsters, projectiles, effects, and profile events.", "The active screen draws the world, entities, player, HUD, and overlays.", "`pg.display.flip()` presents the completed frame."]],
        p("5.3 Screen Architecture", s["Section"]),
        table([["Screen", "Responsibility"], ["SetupScreen", "Create/select profile and copy snapshot into shared game state"], ["MenuScreen", "Start a run, open inventory/profile, or exit"], ["GameScreen", "Generate world, run gameplay, interactions, combat, and outcomes"], ["InventoryScreen", "Show and use persistent/run items"], ["PauseScreen", "Suspend or resume the active game"], ["ProfileScreen", "Show player identity and progression"], ["MenuInventoryScreen", "Inspect inventory from the menu"]], [1.4 * inch, 4.6 * inch]),
        p("5.4 Internal Data Flow", s["Section"]),
        p("The shared `game_state` dictionary carries player identity, class, inventory, equipped items, progression, temporary run items, kill counts, and run flags between screens. `ClientSession` keeps the durable copy in `profiles.json`. A generated `GeneratedWorld` carries tiles, spawn, door key, door, treasures, keys, scrolls, monsters, decorations, scheduled spawns, and reachable coordinates into `GameScreen`.", s["BodyReport"]),
        p("5.5 Objective and Interaction Flow", s["Section"]),
        p("The placement module finds a farthest reachable tile for the door key, then another farthest tile for the door. Treasures are placed on reachable locations and keys are placed away from their corresponding treasures. During play, the player resolves nearby world interactions, collects items, opens treasures when the required key exists, and changes the door from locked to open when the door key condition is satisfied.", s["BodyReport"]),
        PageBreak(),
        p("6 AI and Procedural Generation", s["Chapter"]),
        p("6.1 AI Definition in This Project", s["Section"]),
        p("The project does not train a neural network. Its AI is an explainable algorithmic system that makes generation and difficulty decisions from rules, search results, metrics, randomness, and player history. This is appropriate for a procedural game because the desired behavior is controllable, testable, and fast enough to run when a level starts.", s["BodyReport"]),
        p("6.2 World Generation Pipeline", s["Section"]),
        table([["Stage", "Algorithm / module", "Purpose"], ["1", "genetic_algorithm.generate_population", "Create 8 random 81 x 81 seed grids"], ["2", "layout.carve_nine_sector_maze", "Create a nine-sector maze and connect sector centers"], ["3", "cellular_automata.smooth", "Smooth local wall/floor patterns"], ["4", "bfs.repair_connectivity", "Carve paths between disconnected walkable regions"], ["5", "bfs.evaluate", "Measure open ratio, reachable count, and reachable ratio"], ["6", "genetic_algorithm.crossover/mutate", "Breed two survivors into the next population"], ["7", "decision_tree.choose_best_candidate", "Select the highest profile-weighted candidate"], ["8", "placement.place_entities", "Place objectives, rewards, monsters, lights, and decoration"]], [0.45 * inch, 2.1 * inch, 3.45 * inch]),
        p("6.3 Genetic Algorithm", s["Section"]),
        p("Each generation evaluates candidate maps. The best two grids survive, a crossover combines regions from both parents, and mutation changes individual tiles with a rate that rises by generation: 0.06, 0.08, 0.10, and 0.12 under the current defaults. Four generations are processed. The algorithm is not optimizing a single visual score: it balances walkability, openness, reachable area, and profile statistics through `score_candidate`.", s["BodyReport"]),
        p("6.4 Search Algorithms", s["Section"]),
        p("BFS starts at the selected spawn and visits four-directional walkable neighbors. It supplies reachability and distance maps, finds farthest objective locations, and repairs disconnected regions. A* uses Manhattan distance as its heuristic and checks whether a monster or entity has a valid route to the player. These checks prevent important content from being placed in inaccessible regions.", s["BodyReport"]),
        p("6.5 Profile-aware Decision Tree", s["Section"]),
        p("`mine_recent_rules` reads at most the last five games and converts each into a compact transaction containing outcome, performance, and collected items. It estimates two associations: high kills -> clear and low kills -> loss. `monster_limit_for_profile` then applies transparent rules: two or more consecutive losses plus a strong low-kill/loss association increases the next monster limit up to 18; a strong high-kill/clear association increases it modestly; otherwise the default is 12. This creates adaptive difficulty without hiding the reason for the change.", s["BodyReport"]),
        p("6.6 Entity Placement Strategy", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["Objectives use farthest reachable positions to encourage exploration.", "Treasures receive randomized rewards such as coins, bow, or light.", "Keys are placed away from their target treasures, creating navigation pressure.", "Monsters guard treasures and are also scheduled to appear later in the run.", "Wall lights and liquid clusters add environmental variety after functional placement."]],
        PageBreak(),
        p("7 System Development and Implementation", s["Chapter"]),
        p("7.1 Programming Platform", s["Section"]),
        p("Python is used for application logic, game state, algorithms, and persistence. Pygame provides the event loop, display, timing, drawing, image loading, input, and sprite support. JSON stores profiles locally. The application starts at the root `main.py` and imports the client package after adding the source directories to `sys.path`.", s["BodyReport"]),
        p("7.2 Hardware and Software Requirements", s["Section"]),
        table([["Category", "Recommended baseline"], ["Operating system", "Windows, Linux, or macOS"], ["Runtime", "Python 3 with Pygame 2.5.8"], ["Memory", "4 GB RAM or more"], ["Storage", "Project files plus image assets"], ["Display", "A desktop resolution supported by `config.py`"], ["Development", "VS Code recommended"]], [1.55 * inch, 4.45 * inch]),
        p("7.3 Profile and Run Persistence", s["Section"]),
        p("Registration creates a default player record. Selection loads the snapshot into `game_state`. At the beginning of a run, the game stores a baseline profile and inventory. On clear or death, `ClientSession` records the result, keeps only the five most recent games for the adaptive rules, removes temporary keys, and commits persistent items, coins, and item usage.", s["BodyReport"]),
        p("7.4 Important Implementation Note", s["Section"]),
        p("The repository also contains an older `src/shared/tilemap.py` implementation and a separate `demo/` pipeline. The running client imports `shared.mapgen.generate_map`, so this report treats `src/shared/mapgen/` as the authoritative active implementation and describes the other path as legacy/demo material.", s["BodyReport"]),
        PageBreak(),
        p("8 Testing and Debugging", s["Chapter"]),
        p("8.1 Testing Approach", s["Section"]),
        p("Testing should combine module checks, deterministic seeded generation, gameplay smoke tests, and persistence checks. A fixed seed makes a generated map reproducible while random seeds verify variety. The most valuable assertions are connectivity, reachable objectives, valid monster paths, and correct run-history updates.", s["BodyReport"]),
        p("8.2 Test Cases", s["Section"]),
        table([["ID", "Test", "Expected result"], ["T1", "Launch root `main.py`", "Pygame window opens at SetupScreen"], ["T2", "Create explorer with a name", "Profile is saved and menu opens"], ["T3", "Select existing explorer", "Snapshot restores identity, items, and profile"], ["T4", "Start game with fixed seed", "World is generated with a valid spawn"], ["T5", "Inspect reachable set", "Important regions are connected after repair"], ["T6", "Check door key and door", "Both are reachable and placed separately"], ["T7", "Move into walls and liquids", "Walls block; terrain changes speed/effects"], ["T8", "Attack a monster", "Projectile and health/counter logic update"], ["T9", "Wait for scheduled spawn", "Monster enters active list at its spawn time"], ["T10", "Clear or lose a run", "Outcome is recorded and temporary keys are removed"], ["T11", "Use recent losses", "Next map monster limit follows decision rules"], ["T12", "Run syntax compilation", "Python source compiles without syntax errors"]], [0.45 * inch, 2.55 * inch, 3.0 * inch]),
        p("8.3 Current Verification", s["Section"]),
        p("The deliverable generator itself is executable in the configured project virtual environment. The project source can be syntax-compiled independently. A full interactive playtest requires the Pygame dependency and a graphical session; this report does not claim a completed GUI playthrough.", s["BodyReport"]),
        PageBreak(),
        p("9 Conclusion and Future Work", s["Chapter"]),
        p("Dungeon Exploration combines a clear game loop with a technically interesting procedural generation system. The central design strength is the layered AI pipeline: generation produces candidates, search algorithms enforce playability, placement creates navigable objectives, and recent-history rules adapt enemy quantity. The screen architecture and local profile store provide a practical foundation for continued gameplay work.", s["BodyReport"]),
        p("Future improvements", s["Section"]),
        *[bullet(x, s["BulletReport"]) for x in ["Add automated tests for generator invariants and decision-tree boundary cases.", "Expose a debug seed and generation metrics in a developer overlay.", "Improve monster behavior with line-of-sight, pursuit states, attack range, and difficulty parameters.", "Add audio, richer class-specific abilities, more level objectives, and boss encounters.", "Separate gameplay services from rendering further and remove legacy duplicate map-generation code.", "Add optional multiplayer only after the single-player rules and persistence model are stable."]],
        p("References", s["Chapter"]),
        *[bullet(x, s["BulletReport"]) for x in ["Project source: `src/client/main.py`, `src/client/session.py`, `src/client/screens/game_screen.py`.", "Project source: `src/shared/mapgen/generator.py`, `genetic_algorithm.py`, `layout.py`, `bfs.py`, `astar.py`, `decision_tree.py`, and `placement.py`.", "Supplied reference: `documentation/documentation_sample.docx`.", "Supplied reference: `documentation/dungeon_exploration.docx`.", "Supplied reference: `documentation/dungeon_quest (Repaired).docx`.", "Pygame documentation: https://www.pygame.org/docs/."]],
    ]
    return story


def build_pdf():
    output = OUT_DIR / "Dungeon_Exploration_Documentation.pdf"
    frame = Frame(0.65 * inch, 0.72 * inch, A4[0] - 1.3 * inch, A4[1] - 1.35 * inch, id="normal")
    doc = BaseDocTemplate(str(output), pagesize=A4, rightMargin=0.65 * inch, leftMargin=0.65 * inch, topMargin=0.65 * inch, bottomMargin=0.72 * inch, title="Dungeon Exploration Documentation", author="Dungeon Exploration Project")
    doc.addPageTemplates([PageTemplate(id="report", frames=frame, onPage=header_footer)])
    doc.build(report_story())
    return output


NAVY = RGBColor(22, 50, 79)
TEAL = RGBColor(35, 122, 119)
INK = RGBColor(36, 49, 58)
PALE = RGBColor(242, 247, 248)


def add_text(slide, text, x, y, w, h, size=22, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.42, 12.0, 0.55, size=27, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.05, 11.7, 0.45, size=12, color=TEAL)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.45), Inches(1.0), Inches(0.08))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.fill.background()


def add_bullets(slide, items, x=0.9, y=1.8, w=11.4, size=19, gap=0.52):
    for index, item in enumerate(items):
        add_text(slide, "- " + item, x, y + index * gap, w, 0.4, size=size, color=INK)


def add_flow(slide, labels, y=3.0, x=0.6, width=12.1):
    gap = 0.12
    box_w = (width - gap * (len(labels) - 1)) / len(labels)
    for index, label in enumerate(labels):
        left = x + index * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(box_w), Inches(0.95))
        shape.fill.solid(); shape.fill.fore_color.rgb = PALE
        shape.line.color.rgb = TEAL
        shape.text_frame.text = label
        shape.text_frame.word_wrap = True
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = "Aptos"; paragraph.font.size = Pt(13); paragraph.font.bold = True; paragraph.font.color.rgb = NAVY
        if index < len(labels) - 1:
            add_text(slide, ">", left + box_w, y + 0.27, gap, 0.3, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PALE
    add_text(slide, "Dungeon Exploration", 0.8, 1.55, 11.8, 0.9, size=42, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Procedural action-adventure game", 0.8, 2.55, 11.8, 0.5, size=22, color=TEAL, align=PP_ALIGN.CENTER)
    add_text(slide, "Project presentation | Python + Pygame + algorithmic AI", 0.8, 5.75, 11.8, 0.4, size=15, color=INK, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "Project idea", "A replayable dungeon loop with explainable adaptive difficulty")
    add_bullets(slide, ["Create or select an explorer profile.", "Enter a newly generated dungeon on every run.", "Explore, fight, collect keys and treasures, and unlock the door.", "Persist progression and use recent outcomes to shape future challenge."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Overall game flow", "From launch to a completed or failed run")
    add_flow(slide, ["Launch", "Profile", "Menu", "Generate world", "Explore + fight", "Clear / death"])
    add_text(slide, "The active screen owns input, update, and drawing while shared game_state carries the run across screens.", 1.0, 5.2, 11.2, 0.6, size=18, color=INK, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "System architecture", "Pygame presentation, shared game state, local persistence, and mapgen services")
    add_flow(slide, ["Pygame loop", "ScreenManager", "GameScreen", "shared.mapgen", "ClientSession / JSON"], y=2.55)
    add_bullets(slide, ["One active screen receives events and renders each frame.", "GameScreen bridges generated data with player, monsters, projectiles, and interactions.", "ClientSession saves persistent inventory, coins, item usage, and recent results."], x=1.0, y=4.65, size=17, gap=0.55)

    slide = prs.slides.add_slide(blank); add_title(slide, "Procedural generation pipeline", "A hybrid of evolutionary search, cellular smoothing, and graph search")
    add_flow(slide, ["Seed population", "Carve 9 sectors", "Smooth", "Repair BFS", "Evaluate", "Breed", "Select + place"] , y=2.55)
    add_text(slide, "Default configuration: 81 x 81 map, 8 candidates, 4 generations, 3 x 3 sector structure.", 1.0, 5.25, 11.3, 0.5, size=18, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "Where the AI is", "AI here means explainable algorithmic decisions, not a trained neural network")
    add_bullets(slide, ["Genetic algorithm: keeps strong layouts, crosses parents, and mutates tiles.", "BFS: measures reachable space, distances, and connectivity repairs.", "A*: validates paths before placing or scheduling entities.", "Decision rules: mine the last five games and adjust the next monster limit."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Adaptive difficulty", "The decision tree turns recent player outcomes into a transparent rule")
    add_flow(slide, ["Last 5 games", "Kill + clear data", "Association rules", "Loss streak", "Monster limit"] , y=2.45)
    add_text(slide, "Default = 12 monsters. Repeated low-kill losses can increase the limit up to 18; repeated high-kill clears increase it modestly.", 1.0, 5.1, 11.3, 0.7, size=18, color=INK, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "Objective and entity placement", "Functional content is placed after the map is proven playable")
    add_bullets(slide, ["Door key and door use farthest reachable tiles to create exploration goals.", "Treasures receive randomized rewards: coins, bow, or light.", "Keys are separated from treasure locations.", "Monsters guard treasure and appear through a spawn schedule.", "Lights, water, lava, and decorations add environmental identity."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Gameplay loop", "Real-time control inside the generated grid")
    add_flow(slide, ["Input", "Collision", "Terrain", "Combat", "Items", "Outcome"] , y=2.5)
    add_bullets(slide, ["WASD / arrows move; Shift sprints while stamina is available.", "Mouse input aims and fires projectiles.", "Walls block movement; water and hazards modify survival and movement.", "E or Space opens inventory / resolves world interactions; Escape pauses."] , x=1.0, y=4.65, size=16, gap=0.48)

    slide = prs.slides.add_slide(blank); add_title(slide, "Persistence and progression", "The local profile is part of the game design")
    add_bullets(slide, ["Profiles are stored in `src/client/data/profiles.json`.", "Class, level, experience, inventory, coins, and item usage survive between runs.", "Temporary keys and run-only items are removed when a run ends.", "The last five game outcomes feed the next generation's difficulty decision."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Testing strategy", "What should be demonstrated and verified")
    add_bullets(slide, ["Launch and profile creation / selection.", "Deterministic seeded map generation and random-seed variety.", "BFS reachable ratio and A* path checks for objectives and monsters.", "Movement, collision, combat, item use, scheduled spawning, and door interaction.", "Run recording, persistence, and adaptive monster-limit boundaries."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Conclusion", "A compact game with a strong algorithmic core")
    add_bullets(slide, ["The game combines a readable player loop with fresh procedural worlds.", "The AI is explainable, testable, and directly connected to gameplay variety.", "The architecture leaves room for stronger monster behavior, more objectives, and automated tests.", "The current implementation is local single-player; multiplayer and ML are future possibilities, not present features."])

    slide = prs.slides.add_slide(blank); add_title(slide, "Thank you", "Dungeon Exploration")
    add_text(slide, "Questions?", 0.8, 2.45, 11.8, 0.8, size=38, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "The key takeaway: every run is generated, checked, and tuned before the player enters it.", 1.2, 4.0, 11.0, 0.7, size=21, color=NAVY, align=PP_ALIGN.CENTER)

    output = OUT_DIR / "Dungeon_Exploration_Presentation.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    pdf = build_pdf()
    ppt = build_ppt()
    print(pdf)
    print(ppt)